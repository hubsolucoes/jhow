"""Executor do catálogo: encontra o endpoint certo, consulta a API do cliente e gera a planilha.

É o protótipo do que o assistente faz no chat:
    1) buscar   — dada a pergunta do cliente, lista endpoints de LEITURA que a atendem
    2) detalhar — mostra o que será consultado (para o cliente aprovar)
    3) puxar    — executa a consulta e grava um .xlsx com os dados

Uso (a partir da raiz do catálogo):
    .venv\\Scripts\\python scripts\\executar.py buscar "cobranças pagas em agosto" [--sistema asaas]
    .venv\\Scripts\\python scripts\\executar.py detalhar asaas.cobranca.listar
    .venv\\Scripts\\python scripts\\executar.py puxar asaas.cobranca.listar --filtro status=RECEIVED \\
        --credenciais credenciais/asaas.json --ambiente sandbox --max 500 --saida saida/cobrancas.xlsx

Segurança:
    - Só executa endpoints com execucao.seguro_para_executar = true (leitura).
    - Credenciais vêm de um arquivo JSON fora do catálogo versionado (ou de variáveis de ambiente);
      nunca são impressas, gravadas na planilha ou registradas em log.
    - No produto, esse arquivo é substituído pelo cofre do backend.
"""
import argparse
import json
import os
import re
import sys
import time
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
PAUSA_ENTRE_PAGINAS = 0.35  # segundos; respeita limites conservadores
TIMEOUT = 60
PARADAS = {"de", "da", "do", "das", "dos", "em", "no", "na", "nos", "nas", "a", "o", "as", "os",
           "e", "ou", "um", "uma", "para", "por", "com", "que", "quero", "queria", "preciso",
           "gostaria", "me", "meu", "minha", "meus", "minhas", "todos", "todas", "lista", "quais"}


# ---------------------------------------------------------------- utilidades

def normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto.lower())
    return "".join(c for c in texto if not unicodedata.combining(c))


def termos(texto: str) -> set:
    return {t for t in re.findall(r"[a-z0-9_]+", normalizar(texto)) if len(t) > 2 and t not in PARADAS}


def sistemas_disponiveis() -> list:
    return sorted(p.name for p in (RAIZ / "sistemas").iterdir()
                  if p.is_dir() and (p / "endpoints.json").is_file())


def carregar(slug: str):
    pasta = RAIZ / "sistemas" / slug
    sistema = json.loads((pasta / "system.json").read_text(encoding="utf-8"))
    endpoints = json.loads((pasta / "endpoints.json").read_text(encoding="utf-8"))["endpoints"]
    return sistema, endpoints


def achar_endpoint(endpoint_id: str):
    slug = endpoint_id.split(".")[0]
    candidatos = [s for s in sistemas_disponiveis() if endpoint_id.startswith(s + ".")]
    slug = max(candidatos, key=len) if candidatos else slug
    sistema, endpoints = carregar(slug)
    for e in endpoints:
        if e["id"] == endpoint_id:
            return sistema, e
    raise SystemExit(f"Endpoint não encontrado no catálogo: {endpoint_id}")


def caminho(dado, expr):
    """Lê 'a.b.c' em dicionários aninhados; devolve None se o caminho não existir."""
    if expr in (None, "", [], "[]"):  # "[]" = array na raiz da resposta
        return dado
    atual = dado
    for parte in str(expr).split("."):
        if isinstance(atual, dict):
            atual = atual.get(parte)
        elif isinstance(atual, list) and parte.isdigit():
            atual = atual[int(parte)] if int(parte) < len(atual) else None
        else:
            return None
    return atual


# ---------------------------------------------------------------- 1) buscar

def buscar(args):
    consulta = termos(args.pergunta)
    resultados = []
    for slug in ([args.sistema] if args.sistema else sistemas_disponiveis()):
        sistema, endpoints = carregar(slug)
        for e in endpoints:
            ex = e.get("execucao") or {}
            if not ex.get("seguro_para_executar"):
                continue
            texto = " ".join(str(e.get(k, "")) for k in
                             ("id", "nome_fornecedor", "descricao", "quando_usar", "entidade_canonica", "path"))
            alvo = termos(texto)
            pontos = len(consulta & alvo)
            for c in ex.get("colunas_sugeridas") or []:
                if termos(str(c.get("titulo", ""))) & consulta:
                    pontos += 0.5
            if pontos:
                resultados.append((pontos, slug, e))
    resultados.sort(key=lambda r: -r[0])
    if not resultados:
        print("Nenhum endpoint de leitura encontrado para essa pergunta.")
        print("Sistemas no catálogo:", ", ".join(sistemas_disponiveis()))
        return
    print(f'Pergunta: "{args.pergunta}"\n')
    for pontos, slug, e in resultados[:args.limite]:
        print(f"[{pontos:g}] {e['id']}  ({slug})")
        print(f"     {e['metodo']} {e['path']} — {e.get('descricao', '')[:110]}")
        filtros = [f["parametro"] for f in (e["execucao"].get("filtros_recomendados") or [])]
        if filtros:
            print(f"     filtros: {', '.join(filtros[:8])}")
    print("\nPara ver os detalhes e aprovar: executar.py detalhar <id>")


# -------------------------------------------------------------- 2) detalhar

def detalhar(args):
    sistema, e = achar_endpoint(args.endpoint)
    ex = e["execucao"]
    print(f"Sistema:   {sistema['nome']} ({sistema['slug']})")
    print(f"Endpoint:  {e['id']}")
    print(f"Chamada:   {e['metodo']} {e['path']}")
    print(f"O que faz: {e.get('descricao', '')}")
    print(f"Quando:    {e.get('quando_usar', '')}")
    print(f"Executável: {'sim (leitura)' if ex.get('seguro_para_executar') else 'NÃO — ' + str(ex.get('motivo_inseguro'))}")
    pag = ex.get("paginacao") or {}
    print(f"Paginação: {pag.get('tipo')} ({pag.get('param_pagina')}/{pag.get('param_tamanho')}, máx {pag.get('tamanho_max')})")
    print("Filtros recomendados:")
    for f in ex.get("filtros_recomendados") or []:
        print(f"  - {f.get('parametro')}: {f.get('descricao', '')}")
    print("Colunas da planilha:")
    for c in ex.get("colunas_sugeridas") or []:
        print(f"  - {c.get('titulo')} ({c.get('caminho')})")
    cred = (sistema.get("execucao_auth") or {}).get("credenciais_necessarias") or []
    print("Credenciais necessárias:", ", ".join(c["nome"] for c in cred) or "nenhuma")
    print("Fonte:", e.get("fonte_url"))


# ------------------------------------------------------------ autenticação

def preparar_sessao(sistema, credenciais, ambiente):
    """Devolve (sessão configurada, base_url). Nunca imprime valores de credencial."""
    auth = sistema.get("execucao_auth") or {}
    tipo = auth.get("tipo")
    faltando = [c["nome"] for c in auth.get("credenciais_necessarias") or []
                if c.get("obrigatoria", True) and not credenciais.get(c["nome"])]
    if faltando:
        raise SystemExit(f"Credenciais ausentes para {sistema['slug']}: {', '.join(faltando)}")

    s = requests.Session()
    s.headers.update({"Accept": "application/json", "User-Agent": "catalogo-apis-br/0.1"})
    s.headers.update(auth.get("headers_fixos") or {})

    if tipo == "header_api_key":
        s.headers[auth["header"]] = (auth.get("prefixo") or "") + credenciais[auth["credenciais_necessarias"][0]["nome"]]
    elif tipo == "bearer":
        chave = "access_token" if "access_token" in credenciais else auth["credenciais_necessarias"][0]["nome"]
        s.headers["Authorization"] = "Bearer " + credenciais[chave]
    elif tipo == "basic":
        nome = auth["credenciais_necessarias"][0]["nome"]
        s.auth = (credenciais[nome], "") if auth.get("usuario_basic_e_a_credencial") else \
                 (credenciais.get("usuario", ""), credenciais[nome])
    elif tipo in ("oauth2_client_credentials", "oauth2_refresh_token", "mtls_oauth2"):
        if tipo == "mtls_oauth2":
            s.cert = (credenciais["certificado_pem"], credenciais["chave_privada_pem"])
        dados = ({"grant_type": "client_credentials"} if tipo != "oauth2_refresh_token" else
                 {"grant_type": "refresh_token", "refresh_token": credenciais["refresh_token"]})
        r = s.post(auth["token_url"], data=dados, timeout=TIMEOUT,
                   auth=(credenciais["client_id"], credenciais["client_secret"]))
        if r.status_code >= 400:
            raise SystemExit(f"Falha ao obter token: HTTP {r.status_code} — {r.text[:200]}")
        token = r.json().get("access_token")
        if not token:
            raise SystemExit("A resposta do token não trouxe access_token.")
        s.headers["Authorization"] = "Bearer " + token
        if tipo == "oauth2_refresh_token":
            print("  (token renovado; guarde o novo refresh_token se a API tiver rotacionado)")
    elif tipo != "nenhum":
        raise SystemExit(f"Tipo de autenticação não suportado pelo executor: {tipo}")

    base = (auth.get("base_url_por_ambiente") or {}).get(ambiente)
    if not base:
        base = (sistema["api"]["base_urls"] or {}).get(ambiente)
    if not base:
        raise SystemExit(f"Sem base_url para o ambiente '{ambiente}' em {sistema['slug']}.")
    return s, base.rstrip("/")


# ----------------------------------------------------------------- 3) puxar

def coletar(sessao, base, endpoint, filtros, maximo, corpo=None):
    ex = endpoint["execucao"]
    pag = ex.get("paginacao") or {}
    tipo = pag.get("tipo", "nenhuma")
    tamanho = min(pag.get("tamanho_max") or 100, maximo)
    url = base + endpoint["path"]
    registros, pagina, cursor, chamadas = [], 0, None, 0
    MAX_PAGINAS = 500

    while True:
        if pagina >= MAX_PAGINAS:
            print(f"  parando por segurança em {MAX_PAGINAS} páginas")
            break
        params = dict(filtros)
        if tipo == "offset":
            params[pag["param_pagina"]] = len(registros)
            params[pag["param_tamanho"]] = tamanho
        elif tipo == "page":
            params[pag["param_pagina"]] = pagina + 1
            params[pag["param_tamanho"]] = tamanho
        elif tipo in ("cursor", "versao"):
            if cursor is not None:
                params[pag.get("param_cursor") or "cursor"] = cursor
            elif pag.get("param_tamanho"):
                params[pag["param_tamanho"]] = tamanho

        metodo = endpoint["metodo"].upper()
        if metodo == "GET":
            r = sessao.get(url, params=params, timeout=TIMEOUT)
        else:
            # leitura que exige POST (ex.: decodificar QR Code): uma chamada só, sem paginar
            r = sessao.request(metodo, url, params=filtros or None, json=corpo, timeout=TIMEOUT)
        chamadas += 1
        if r.status_code == 429:
            espera = int(r.headers.get("Retry-After", 5))
            print(f"  limite de requisições atingido; aguardando {espera}s")
            time.sleep(espera)
            continue
        if r.status_code >= 400:
            raise SystemExit(f"A API respondeu HTTP {r.status_code}: {r.text[:300]}")
        corpo = r.json()
        lote = caminho(corpo, ex.get("lista_em"))
        if lote is None:
            lote = corpo if isinstance(corpo, list) else [corpo]
        if isinstance(lote, dict):
            lote = [lote]
        registros.extend(lote)
        print(f"  página {pagina + 1}: {len(lote)} registros (total {len(registros)})")

        if endpoint["metodo"].upper() != "GET":
            break
        if tipo == "nenhuma" or len(registros) >= maximo or not lote:
            break
        fim = pag.get("fim")
        if fim == "pagina_incompleta" and len(lote) < tamanho:
            break
        if fim == "hasMore_false" and not caminho(corpo, "hasMore"):
            break
        if tipo == "cursor":
            cursor = caminho(corpo, pag.get("campo_proximo") or "next")
            if not cursor:
                break
        if tipo not in ("offset", "page", "cursor", "versao"):
            print(f"  paginação '{tipo}' não implementada pelo executor; trazendo só a primeira página")
            break
        if tipo == "versao":
            # a próxima página vem com "maior valor desta página" (ex.: versao na Focus NFe)
            campo = pag.get("campo_proximo") or "versao"
            valores = [caminho(reg, campo) for reg in lote]
            valores = [v for v in valores if v is not None]
            if not valores:
                break
            cursor = max(valores)
        pagina += 1
        time.sleep(PAUSA_ENTRE_PAGINAS)

    return registros[:maximo], chamadas


def gerar_planilha(destino: Path, sistema, endpoint, registros, filtros, chamadas):
    ex = endpoint["execucao"]
    colunas = ex.get("colunas_sugeridas") or []
    if not colunas:  # sem colunas sugeridas, usa as chaves do primeiro registro
        chaves = list(registros[0].keys()) if registros and isinstance(registros[0], dict) else ["valor"]
        colunas = [{"titulo": k, "caminho": k, "tipo": "string"} for k in chaves]

    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.append([c["titulo"] for c in colunas])
    for celula in ws[1]:
        celula.font = Font(bold=True)
        celula.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for reg in registros:
        linha = []
        for c in colunas:
            valor = caminho(reg, c["caminho"])
            if isinstance(valor, (dict, list)):
                valor = json.dumps(valor, ensure_ascii=False)
            linha.append(valor)
        ws.append(linha)
    for i, c in enumerate(colunas, start=1):
        largura = max(len(str(c["titulo"])) + 2, 12)
        ws.column_dimensions[get_column_letter(i)].width = min(largura, 42)
    if registros:
        ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Informações")
    for chave, valor in [
        ("Sistema", sistema["nome"]),
        ("Endpoint", endpoint["id"]),
        ("Chamada", f"{endpoint['metodo']} {endpoint['path']}"),
        ("Descrição", endpoint.get("descricao", "")),
        ("Filtros aplicados", json.dumps(filtros, ensure_ascii=False) if filtros else "nenhum"),
        ("Registros", len(registros)),
        ("Requisições feitas", chamadas),
        ("Consultado em", datetime.now(timezone.utc).astimezone().strftime("%d/%m/%Y %H:%M")),
        ("Fonte da documentação", endpoint.get("fonte_url", "")),
        ("Observação", "Dados lidos pela API do próprio cliente. Nenhuma credencial é gravada neste arquivo."),
    ]:
        info.append([chave, str(valor)])
    for celula in info["A"]:
        celula.font = Font(bold=True)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 90

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def puxar(args):
    sistema, endpoint = achar_endpoint(args.endpoint)
    ex = endpoint.get("execucao") or {}
    if not ex.get("seguro_para_executar"):
        raise SystemExit(
            f"RECUSADO: {endpoint['id']} não é somente leitura.\n"
            f"Motivo: {ex.get('motivo_inseguro') or 'ação de escrita'}\n"
            "O assistente só executa consultas; para escrever, entregue o código ao cliente."
        )
    filtros = dict(p.split("=", 1) for p in args.filtro) if args.filtro else {}
    corpo = json.loads(args.corpo) if args.corpo else None
    if endpoint["metodo"].upper() != "GET" and corpo is None:
        raise SystemExit(
            f"{endpoint['id']} é leitura, mas usa {endpoint['metodo']}: "
            "informe o corpo da requisição.\n"
            "Exemplo: --corpo '{\"payload\": \"<<CONTEUDO>>\"}'"
        )
    if args.simular:
        print(f"Simulação: {endpoint['metodo']} {endpoint['path']} com filtros {filtros or 'nenhum'}")
        return

    credenciais = {}
    arq = Path(args.credenciais) if args.credenciais else RAIZ / "credenciais" / f"{sistema['slug']}.json"
    if arq.is_file():
        credenciais = json.loads(arq.read_text(encoding="utf-8"))
        print(f"Credenciais lidas de {arq} ({len(credenciais)} campos)")
    for c in (sistema.get("execucao_auth") or {}).get("credenciais_necessarias") or []:
        ambiente_var = os.environ.get(f"{sistema['slug'].upper().replace('-', '_')}_{c['nome'].upper()}")
        if ambiente_var:
            credenciais[c["nome"]] = ambiente_var

    sessao, base = preparar_sessao(sistema, credenciais, args.ambiente)
    print(f"Consultando {sistema['nome']} ({args.ambiente}) — {endpoint['metodo']} {endpoint['path']}")
    registros, chamadas = coletar(sessao, base, endpoint, filtros, args.max, corpo)
    destino = Path(args.saida) if args.saida else RAIZ / "saida" / f"{endpoint['id'].replace('.', '_')}.xlsx"
    gerar_planilha(destino, sistema, endpoint, registros, filtros, chamadas)
    print(f"\n{len(registros)} registros em {chamadas} requisição(ões).")
    print(f"Planilha: {destino}")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="comando", required=True)

    b = sub.add_parser("buscar", help="encontra endpoints de leitura para a pergunta do cliente")
    b.add_argument("pergunta")
    b.add_argument("--sistema")
    b.add_argument("--limite", type=int, default=5)
    b.set_defaults(func=buscar)

    d = sub.add_parser("detalhar", help="mostra o que será consultado, para o cliente aprovar")
    d.add_argument("endpoint")
    d.set_defaults(func=detalhar)

    x = sub.add_parser("puxar", help="executa a consulta e grava a planilha")
    x.add_argument("endpoint")
    x.add_argument("--filtro", action="append", metavar="chave=valor")
    x.add_argument("--credenciais")
    x.add_argument("--ambiente", default="sandbox")
    x.add_argument("--max", type=int, default=1000)
    x.add_argument("--saida")
    x.add_argument("--corpo", metavar="JSON", help="corpo da requisição, para leituras que usam POST")
    x.add_argument("--simular", action="store_true")
    x.set_defaults(func=puxar)

    args = p.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
